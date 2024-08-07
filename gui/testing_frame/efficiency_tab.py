import tkinter as tk
from tkinter import ttk, messagebox
import threading
import queue
import sys
import time
import numpy as np
from io import StringIO
from ..setting_frame.EfficiencyTest import EfficiencyTestFrame
from .measure_efficiency import *
from config import *
from .test_tab import *
import datetime
import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd
from pandastable import Table, TableModel
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import warnings
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
import pandas as pd
from pandastable import Table, TableModel
import os  # Add this line
import xlsxwriter  # Make sure this is also imported

class EfficiencyTab(TestTab):
    def __init__(self, parent, instrument_manager, setting_frame):
        super().__init__(parent, instrument_manager, setting_frame, "Efficiency")
        self.results = []

    def create_widgets(self):
        self.create_control_frame()
        self.create_results_area()


    def create_control_frame(self):
        control_frame = tk.Frame(self, bd=2, relief=tk.RIDGE)
        control_frame.pack(pady=10, padx=10, fill=tk.X)

        self.start_button = tk.Button(control_frame, text="Start Efficiency Test", command=self.start_efficiency_test)
        self.start_button.pack(side=tk.LEFT, padx=5)

        self.stop_button = tk.Button(control_frame, text="Stop Test", command=self.stop_efficiency_test, state=tk.DISABLED)
        self.stop_button.pack(side=tk.LEFT, padx=5)

        self.progress_var = tk.DoubleVar()
        self.progress_bar = ttk.Progressbar(control_frame, variable=self.progress_var, maximum=100)
        self.progress_bar.pack(side=tk.LEFT, padx=5, expand=True, fill=tk.X)
        
    def create_results_area(self):
        results_frame = tk.Frame(self, bd=2, relief=tk.RIDGE)
        results_frame.pack(pady=10, padx=10, fill=tk.BOTH, expand=True)

        # Add a title for the results area
        self.results_title = tk.Label(results_frame, text="Efficiency Test", font=("times new roman", 16, "bold"), bg='black', fg="white")
        self.results_title.pack(fill=tk.X, padx=5, pady=5)

        # Create a notebook for different views
        self.notebook = ttk.Notebook(results_frame)
        self.notebook.pack(fill=tk.BOTH, expand=True, pady=10, padx=10)

        # Text widget for log output with scrollbar
        log_frame = tk.Frame(self.notebook)
        self.results_text = tk.Text(log_frame, height=30, width=40)
        scrollbar = tk.Scrollbar(log_frame, command=self.results_text.yview)
        self.results_text.configure(yscrollcommand=scrollbar.set)
        self.results_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.notebook.add(log_frame, text="Log")

        # Frame for Excel data
        self.excel_frame = tk.Frame(self.notebook)
        self.notebook.add(self.excel_frame, text="Data Table")
        

        
    def start_efficiency_test(self):
        if self.test_running:
            messagebox.showwarning("Warning", "Test is already running")
            return

        try:
            validated_settings = self.setting_frame.current_test_frame.get_values()
            if not validated_settings:
                messagebox.showerror("Error", "No validated settings available. Please set values in the Setting Frame.")
                return
            
            self.test_running = True
            self.start_button.config(state=tk.DISABLED)
            self.stop_button.config(state=tk.NORMAL)

            self.redirect_output()
            self.test_thread = threading.Thread(target=self.run_test_thread, args=(validated_settings,))
            self.test_thread.start()
            reset_stop_flag()
            self.after(100, self.check_test_thread)
            self.after(100, self.update_output)

        except Exception as e:
            messagebox.showerror("Error", f"An error occurred: {str(e)}")
            self.test_running = False
            self.start_button.config(state=tk.NORMAL)
            self.stop_button.config(state=tk.DISABLED)
            
            
    def run_test_thread(self, validated_settings):
        sys.stderr.write("Debug: run_test_thread started\n")
        sys.stderr.flush()
        try:
            self.setup_progress_bar(validated_settings)
            self.start_time = time.time()
            self.after(0, self.update_progress_by_time)
            stop_flag = False
            self.print_validated_settings(validated_settings)
            
            # Run the test and get results
            try:
                sys.stderr.write("Debug: About to call eff function\n")
                sys.stderr.flush()
                self.results = efficiency(**validated_settings,instrument_manager=self.instrument_manager)
                sys.stderr.write("Debug: eff function completed\n")
                sys.stderr.flush()
            except Exception as e:
                error_message = f"Error in eff function: {str(e)}\n"
                error_message += "Arguments passed to eff:\n"
                for key, value in validated_settings.items():
                    error_message += f"{key}: {value}\n"
                print(error_message)
                self.update_results(error_message)
                raise

            self.update_results(f"Stop flag set to: {get_stop_flag()}")
            if get_stop_flag():
                self.update_results("Test stopped!")
            else:
                self.update_results("Test completed successfully!")
                self.create_excel_file()
        except Exception as e:
            error_message = f"Error during test: {str(e)}\n"
            error_message += "Traceback:\n"
            import traceback
            error_message += traceback.format_exc()
            print(error_message)
            self.after(0, lambda: self.update_results(error_message))
        finally:
            self.test_running = False
            self.restore_output()
            self.after(0, lambda: self.progress_var.set(100))
                
            
    def redirect_output(self):
        self.old_stdout = sys.stdout
        sys.stdout = StringIO()

    def restore_output(self):
        sys.stdout = self.old_stdout

    def update_output(self):
        if self.test_running:
            output = sys.stdout.getvalue()
            if output:
                self.results_text.insert(tk.END, output)
                self.results_text.see(tk.END)
                sys.stdout.truncate(0)
                sys.stdout.seek(0)
            self.after(100, self.update_output)
            

            
            
    def setup_progress_bar(self, settings):
        input_v_count = len(settings['Input_V'])
        low_load_steps = len(np.arange(settings['Low_load_start'], settings['Low_load_stop'], settings['Low_load_step']))
        high_load_steps = len(np.arange(settings['High_load_start'], settings['High_load_stop'] + 0.002, settings['High_load_step']))
        
        total_steps = input_v_count * (low_load_steps + high_load_steps)
        total_time = (low_load_steps * settings['low_load_timing'] + 
                    high_load_steps * settings['high_load_timing']) * input_v_count
        
        self.total_estimated_time = total_time
        self.progress_update_interval = 0.5  # Update every 0.5 seconds
        
    def update_progress_by_time(self):
        if not hasattr(self, 'start_time'):
            self.start_time = time.time()
        
        elapsed_time = time.time() - self.start_time
        progress = min(elapsed_time / self.total_estimated_time * 100, 100)
        self.progress_var.set(progress)
        
        if progress < 100 and self.test_running:
            self.after(int(self.progress_update_interval * 1000), self.update_progress_by_time)
        else:
            self.progress_var.set(100)
        

    def update_progress(self, step):
        progress = (step / self.total_progress_steps) * 100
        self.progress_var.set(progress)

    def check_if_stopped(self):
        return not self.test_running

    def check_test_thread(self):
        if self.test_thread.is_alive():
            self.after(100, self.check_test_thread)
        else:
            self.start_button.config(state=tk.NORMAL)
            self.stop_button.config(state=tk.DISABLED)

    def stop_efficiency_test(self):
        if not self.test_running:
            return
      
        self.test_running = False
        self.update_results("Test stopped by user.")
        set_stop_flag()
        

        self.start_button.config(state=tk.NORMAL)
        self.stop_button.config(state=tk.DISABLED)
        self.progress_var.set(100)  # Set progress to 100% when stopped


    def update_progress(self, value):
        self.progress_var.set(value)

    def update_results(self, message):
        # Use after() to ensure this runs on the main thread
        self.after(0, lambda: self.results_text.insert(tk.END, message + "\n"))
        self.after(0, lambda: self.results_text.see(tk.END))
        self.after(0, lambda: self.results_text.update_idletasks())  # Force update of the widget

    def lock_frame(self):
        self.start_button.config(state=tk.DISABLED)

    def unlock_frame(self):
        self.start_button.config(state=tk.NORMAL)

    def print_validated_settings(self, validated_settings):
        print_string = f"""
        Input Shunt Parameters:
            - Max Voltage: {validated_settings['input_shunt_max_voltage']}
            - Max Current: {validated_settings['input_shunt_max_current']}

        Output Shunt Parameters:
            - Max Voltage: {validated_settings['output_shunt_max_voltage']}
            - Max Current: {validated_settings['output_shunt_max_current']}

        Channel Assignments:
            - Input Voltage: {validated_settings['input_v_ch']}
            - Input Current: {validated_settings['input_i_ch']}
            - Output Voltage: {validated_settings['output_v_ch']}
            - Output Current: {validated_settings['output_i_ch']}
            - Vcc: {validated_settings['vcc_ch']}
            - LDO: {validated_settings['ldo_ch']}

        Maximum Ratings:
            - Input Voltage: {validated_settings['Max_input_voltage']}
            - Input Current: {validated_settings['Max_input_current']}
            - Load Current: {validated_settings['Max_load_current']}

        Load Sweep Parameters:
            - Input Voltage: {validated_settings['Input_V']}
            - Input Current: {validated_settings['Input_I']}

        Low Load Sweep:
            - Start: {validated_settings['Low_load_start']}
            - Step: {validated_settings['Low_load_step']}
            - Stop: {validated_settings['Low_load_stop']}

        High Load Sweep:
            - Start: {validated_settings['High_load_start']}
            - Step: {validated_settings['High_load_step']}
            - Stop: {validated_settings['High_load_stop']}

        Timing Parameters:
            - Low Load: {validated_settings['low_load_timing']}
            - High Load: {validated_settings['high_load_timing']}

        Frequency: {validated_settings['FRE']}
        """
        self.update_results(print_string)

    def generate_filename(self):
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"efficiency_test_results_{timestamp}.xlsx"   


        
    def display_excel_in_gui(self, filename):
        # Suppress warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            
            # Read the Excel file
            df = pd.read_excel(filename)

            # Clear previous content
            for widget in self.excel_frame.winfo_children():
                widget.destroy()

            # Create a Table widget
            self.table = Table(self.excel_frame, dataframe=df, showtoolbar=True, showstatusbar=True)
            self.table.show()
            

    def create_chart(self, workbook, worksheet, title, x_title, y_title, x_col, y_col, position, headers):
        chart = workbook.add_chart({'type': 'scatter', 'subtype': 'straight_with_markers'})
        chart.set_title({'name': title, 'name_font': {'size': 14, 'bold': True}})
        chart.set_x_axis({'name': x_title, 'major_gridlines': {'visible': True}, 'name_font': {'size': 14, 'bold': True}})
        chart.set_y_axis({'name': y_title, 'name_font': {'size': 14, 'bold': True}})
        chart.set_legend({'position': 'bottom'})

        colors = ['#0000FF', '#FFA500', '#008000', '#FF0000']  # Blue, Orange, Green, Red
        markers = ['circle', 'diamond', 'triangle', 'square']

        row = 1
        for i, (input_voltage, measurements) in enumerate(sorted(self.results.items())):
            end_row = row + len(measurements) - 1
            chart.add_series({
                'name': f'Vin = {input_voltage}V',
                'categories': [worksheet.name, row, x_col, end_row, x_col],
                'values': [worksheet.name, row, y_col, end_row, y_col],
                'marker': {'type': markers[i % len(markers)], 'size': 7},
                'line': {'color': colors[i % len(colors)], 'width': 2.25},
            })
            row = end_row + 1

        chart.set_size({'width': 500, 'height': 300})
        worksheet.insert_chart(position, chart)
        return chart

    def create_efficiency_chart(self, workbook, worksheet, headers):
        return self.create_chart(workbook, worksheet, 'Efficiency Measurement',
                                'Load Current(A)', 'Efficiency (%)',
                                headers.index('electronic_load_current'), headers.index('efficiency'), 'S2', headers)

    def create_load_regulation_chart(self, workbook, worksheet, headers):
        return self.create_chart(workbook, worksheet, 'Load Regulation', 
                                'Load Current(A)', 'Output Voltage(V)', 
                                headers.index('electronic_load_current'), headers.index('v_output_voltage'), 'S20', headers)

    def create_vcc_regulation_chart(self, workbook, worksheet, headers):
        return self.create_chart(workbook, worksheet, 'VCC Regulation', 
                                'Load Current(A)', 'VCC Voltage(V)', 
                                headers.index('electronic_load_current'), headers.index('VCC'), 'S38', headers)

    def create_pg_regulation_chart(self, workbook, worksheet, headers):
        return self.create_chart(workbook, worksheet, 'PG Regulation', 
                                'Load Current(A)', 'PG Voltage(V)', 
                                headers.index('electronic_load_current'), headers.index('PG'), 'S56', headers)

    def create_fsw_regulation_chart(self, workbook, worksheet, headers):
        return self.create_chart(workbook, worksheet, 'FSW Regulation', 
                                'Load Current(A)', 'Frequency(Hz)', 
                                headers.index('electronic_load_current'), headers.index('switching_frequency'), 'S74', headers)

    def create_excel_file(self):
        # Create the results/efficiency directory if it doesn't exist
        result_dir = os.path.join("results", "efficiency")
        os.makedirs(result_dir, exist_ok=True)

        selected_ic = self.setting_frame.current_test_frame.selected_ic
        test_type = self.setting_frame.test_type_combo.get()
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{selected_ic}_{test_type}_{timestamp}.xlsx"
        filepath = os.path.join(result_dir, filename)
        
        workbook = xlsxwriter.Workbook(filepath)
        worksheet = workbook.add_worksheet()

        headers = ['input_voltage_setpoint', 'VCC', 'PG', 'v_input_shunt', 'v_output_shunt',
                    'v_input_voltage', 'v_output_voltage', 'i_input_current', 'i_output_current',
                    'p_input_power', 'p_output_power', 'p_power_loss', 'efficiency',
                    'electronic_load_setpoint', 'electronic_load_current', 'electronic_load_voltage',
                    'switching_frequency']

        # Write headers
        for col, header in enumerate(headers):
            worksheet.write(0, col, header)

        # Write data
        row = 1
        for input_voltage, measurements in self.results.items():
            for measurement in measurements:
                worksheet.write(row, 0, input_voltage)
                for col, header in enumerate(headers[1:], start=1):
                    worksheet.write(row, col, measurement.get(header, ''))
                row += 1

        # Create charts
        self.create_efficiency_chart(workbook, worksheet, headers)
        self.create_load_regulation_chart(workbook, worksheet, headers)
        self.create_vcc_regulation_chart(workbook, worksheet, headers)
        self.create_pg_regulation_chart(workbook, worksheet, headers)
        self.create_fsw_regulation_chart(workbook, worksheet, headers)

        # Adjust column widths
        for i, header in enumerate(headers):
            worksheet.set_column(i, i, 15)  # Set each column width to 15

        workbook.close()
        self.update_results(f"Excel file created: {filepath}")
        self.display_excel_in_gui(filepath)